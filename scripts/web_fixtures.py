"""Synthetic fixtures and the actual desktop privacy protocol for browser QA."""
import io,json,sys,hashlib
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from redaction import redact_markdown,attach_identity,seal_payload,encrypt_mapping
from recovery import build_pack,export_record
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
ROOT=Path(sys.argv[1] if len(sys.argv)>1 else '/tmp/mad-web-fixtures');ROOT.mkdir(parents=True,exist_ok=True)
source='# 中文合同\n\n甲方：测试单位\n电话：13800138000\n邮箱：contract@example.com\n金额：人民币 125000.50 元\n地址：北京市测试路10号\n日期：2026年10月15日\n案号：（2026）京0101民初123号\n统一社会信用代码：91110000MA0000000X\n身份证：110101199001011234\n账户：6222021234567890123\n合同编号：MAD-2026-001\n白名单：13800138001\n'
(ROOT/'source.md').write_text(source,encoding='utf8')
image=Image.new('RGB',(1000,280),'white');draw=ImageDraw.Draw(image)
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',38) if Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf').exists() else ImageFont.load_default(size=38)
draw.text((40,40),'Local OCR sample',fill='black',font=font);draw.text((40,110),'Phone 13800138000',fill='black',font=font);image.save(ROOT/'scan.png')
cjk_paths=[str(Path.home()/'.local/share/fonts/NotoSansCJKsc-Regular.otf'),'/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc','C:/Windows/Fonts/msyh.ttc','/System/Library/Fonts/PingFang.ttc']
cjk=None
for candidate in cjk_paths:
 try:
  if Path(candidate).is_file():
   cjk=candidate
   break
 except OSError:
  # An inaccessible optional font must not fail protocol or CI fixtures.
  continue
if cjk:
 chinese_image=Image.new('RGB',(1000,280),'white');chinese_draw=ImageDraw.Draw(chinese_image);chinese_font=ImageFont.truetype(cjk,42)
 chinese_draw.text((40,35),'中文合同 金额 125000.50 元',fill='black',font=chinese_font)
 chinese_draw.text((40,115),'联系电话 13800138000',fill='black',font=chinese_font);chinese_image.save(ROOT/'chinese-scan.png')

d=Document();d.add_heading('中文合同',0);d.add_paragraph(source);t=d.add_table(rows=2,cols=2);t.cell(0,0).text='项目';t.cell(0,1).text='金额';t.cell(1,0).text='货款';t.cell(1,1).text='125000.50';d.add_picture(str(ROOT/'scan.png'));d.save(ROOT/'中文合同.docx')
w=Workbook();w.active.title='金额';w.active.append(['项目','金额']);w.active.append(['货款',125000.50]);w.create_sheet('联系人').append(['电话','13800138000']);w.save(ROOT/'中文表格.xlsx')
r=Presentation();slide=r.slides.add_slide(r.slide_layouts[1]);slide.shapes.title.text='中文汇报';slide.placeholders[1].text='13800138000\n人民币 125000.50 元';r.save(ROOT/'中文汇报.pptx')
(ROOT/'中文合同.rtf').write_text(r'{\rtf1\ansi\ansicpg936 '+''.join('\\u'+str(ord(c) if ord(c)<32768 else ord(c)-65536)+'?' if ord(c)>127 else c for c in '中文合同 13800138000')+'}',encoding='ascii')
c=Canvas(str(ROOT/'mixed.pdf'));c.drawString(40,750,'Text page 1');c.showPage();c.drawImage(ImageReader(image),40,450,width=500,height=140);c.save()
password='fixture-password';pid=uuid4().hex;did=uuid4().hex
options={'categories':['phone','email','id_card','bank_account','credit_code','case_number','date','amount','address','file_code'],'keywords':['测试单位'],'whitelist':['13800138001']}
masked,mapping,summary=redact_markdown(source,options,pid);masked=attach_identity(masked,did,pid)
payload={'documentId':did,'processingId':pid,'sourceHash':hashlib.sha256(source.encode()).hexdigest(),'redactedHash':hashlib.sha256(masked.encode()).hexdigest(),'mapping':mapping}
vault=seal_payload(payload,password,'MAD-R2');pack=build_pack([vault],password,[export_record('source.md',masked)])
legacy_text,legacy_mapping,_=redact_markdown(source,options)
legacy={'format':'MarkAllDownEncryptedVaultPack','version':1,'vaults':[encrypt_mapping(legacy_mapping,password)]}
(ROOT/'desktop.json').write_text(json.dumps({'source':source,'content':masked,'pack':pack,'password':password,'options':options,'summary':summary,'legacyContent':legacy_text,'legacyPack':legacy},ensure_ascii=False),encoding='utf8')
print(ROOT)
