"""Actual FastAPI and converter: synthetic fixtures, no external AI requests."""
import io
from pathlib import Path
import sys
import tempfile
import unittest
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from fastapi.testclient import TestClient
import server
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas

class RealApplicationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(server.app)
        cls.headers = {'X-API-Token': server.API_TOKEN}

    def test_real_routes_are_authenticated(self):
        for route in ('/api/redact', '/api/restore', '/api/ai/suggest', '/api/privacy/export'):
            with self.subTest(route=route):
                self.assertEqual(self.client.post(route, json={}).status_code, 401)

    def test_actual_conversion_export_restore_and_legacy_export(self):
        d = Document(); d.add_heading('Synthetic contract', 0)
        d.add_paragraph('Contact 13800138000 and abc@example.com')
        docx = io.BytesIO(); d.save(docx)
        w = Workbook(); w.active.append(['Contact', '13800138000', 'abc@example.com'])
        xlsx = io.BytesIO(); w.save(xlsx)
        pdf = io.BytesIO(); c = Canvas(pdf)
        c.drawString(50, 800, 'Contact 13800138000 abc@example.com'); c.save()
        fixtures = [('markdown.md', b'# Contract\n13800138000 abc@example.com'),
                    ('word.docx', docx.getvalue()), ('sheet.xlsx', xlsx.getvalue()), ('pdf.pdf', pdf.getvalue())]
        with tempfile.TemporaryDirectory() as directory:
            for name, raw in fixtures:
                with self.subTest(name=name):
                    response = self.client.post('/api/convert', headers=self.headers, files={'file': (name, raw)})
                    self.assertEqual(response.status_code, 200, response.text)
                    cache_path = response.json()['cachePath']
                    cached = self.client.get('/api/text', headers=self.headers, params={'path': cache_path})
                    self.assertEqual(cached.status_code, 200, cached.text)
                    original = cached.json()['content']
                    self.assertIn('13800138000', original)
                    redaction = self.client.post('/api/redact', headers=self.headers, json={
                        'content': original, 'passphrase': 'test secure password', 'options': {}})
                    self.assertEqual(redaction.status_code, 200, redaction.text)
                    result = redaction.json()
                    self.assertNotIn('13800138000', result['content'])
                    exported = self.client.post('/api/privacy/export', headers=self.headers, json={
                        'files': [{'name': 'document_0001.md', 'relativePath': 'document_0001.md',
                                   'content': result['content'], 'sourceHash': result['sourceHash'], 'reviewed': True}],
                        'pack': {'format': 'MarkAllDownEncryptedVaultPack', 'version': 1, 'vaults': [result['vault']]},
                        'passphrase': 'test secure password',
                        'options': {'exportPath': directory, 'format': 'individual', 'structure': 'flat'}})
                    self.assertEqual(exported.status_code, 200, exported.text)
                    text = (Path(exported.json()['path']) / 'document_0001.md').read_text(encoding="utf-8")
                    restored = self.client.post('/api/restore', headers=self.headers, json={
                        'content': text, 'vaults': [result['vault']], 'passphrase': 'test secure password'})
                    self.assertEqual(restored.status_code, 200, restored.text)
                    self.assertEqual(restored.json()['content'], original)
                    legacy = self.client.post('/api/export', headers=self.headers, json={
                        'files': [{'name': name, 'cachePath': cache_path}],
                        'options': {'exportPath': directory, 'format': 'individual', 'preserveNames': True, 'structure': 'flat'}})
                    self.assertEqual(legacy.status_code, 200, legacy.text)
                    self.assertEqual((Path(directory) / (Path(name).stem + '.md')).read_text(encoding="utf-8"), original)
